# -*- coding: utf-8 -*-
import scrapy
import openpyxl
from ..items import HudousanItem

class HudousanSpiderSpider(scrapy.Spider):
    name = 'hudousan_spider'
    allowed_domains = ['tancyu-f.com']
    start_urls = ['https://tancyu-f.com/rent/']
    item = HudousanItem()
    # Excelに書き込む際の行番号
    num = 1

    # メイン関数
    def parse(self, response):
        # 不動産のURLをそれぞれ取得
        for url in response.css('div.category__detail a::attr("href")').extract():
            self.item['link'] = url
            yield scrapy.Request(url=url, callback=self.get_data)

        # 次ページのURLを取得する
        next_page = response.css('div.nav_pager a::attr("href")')[-1].extract()
        print(next_page)
        if next_page is not None:
            yield scrapy.Request(url=response.urljoin(next_page), callback=self.parse)

    # 各ページ（物件）の詳細をスクレイピング
    def get_data(self, response):
        self.item['category'] = response.css('div.container em::text')[1].extract()
        self.item['name'] = response.css('div.container h1::text')[0].extract()
        self.item['image'] = response.css('ul.slides img::attr("src")')[-1].extract()
        print(self.item['category'])
        try:
            self.item['price'] = response.css('dl.content-info em::text')[0].extract()
        except:
            self.item['price'] = response.css('dl.content-info dd::text')[0].extract()

        try:
            self.item['operation'] = response.css('dl.content-info dd::text')[1].extract()
            self.item['deposit'] = response.css('dl.content-info dd::text')[2].extract()
            self.item['reward'] = response.css('dl.content-info dd::text')[3].extract()
            self.item['renewal'] = response.css('dl.content-info dd::text')[-3].extract()
            self.item['layout'] = response.css('dl.content-info dd::text')[-2].extract()
            self.item['age'] = response.css('dl.content-info span::text')[-2].extract()
            self.item['space'] = response.css('dl.content-info dd::text')[-1].extract()
        except:
            self.item['operation'] = response.css('dl.content-info dd::text')[1].extract()
            self.item['deposit'] = response.css('dl.content-info dd::text')[2].extract()
            self.item['renewal'] = response.css('dl.content-info dd::text')[-3].extract()
            self.item['layout'] = response.css('dl.content-info dd::text')[-2].extract()
            self.item['age'] = response.css('dl.content-info span::text')[-2].extract()
            self.item['space'] = response.css('dl.content-info dd::text')[-1].extract()
            self.item['reward'] = "なし"

        self.item['access'] = response.css('div.content-catch__inner p::text')[0].extract()
        self.item['location'] = response.css('div.content-map h4::text')[0].extract()
        # 物件概要
        dd = response.css('div.content-detail__frame dd::text')
        self.item['water'] = dd[0].extract()
        self.item['expenses'] = dd[1].extract()
        self.item['parking'] = dd[2].extract()
        self.item['construction'] = dd[3].extract()
        self.item['window'] = dd[4].extract()
        self.item['now'] = dd[5].extract()
        self.item['available'] = dd[6].extract()
        self.item['transaction'] = dd[7].extract()
        self.item['term'] = dd[8].extract()
        self.item['contract'] = dd[9].extract()

        # 特徴・設備
        dd = response.css('div.content-feature dd::text')
        self.item['waterworks'] = dd[0].extract()
        self.item['sewage'] = dd[1].extract()
        self.item['waterline'] = dd[2].extract()
        self.item['gas'] = dd[3].extract()
        self.item['supply'] = dd[5].extract()
        self.item['aircon'] = dd[6].extract()
        self.item['laundry'] = dd[8].extract()
        self.item['dehumidifier'] = dd[10].extract()
        self.item['shower'] = dd[11].extract()
        self.item['toilet'] = dd[12].extract()
        self.item['BS'] = dd[14].extract()
        self.output()

    # Excelファイルに出力
    def output(self):
        # 行番号
        self.num += 1
        # 列項目
        column = {'URL':'link', '画像':'image', '種類':'category',
                  '名前':'name', '賃料':'price', '管理費':'operation',
                  '保証金':'deposit', '礼金':'reward', '区費・町内会費':'expenses',
                  '築年日':'age','場所':'location', 'アクセス':'access',
                  '水道料金':'water','駐車場':'parking', 'トイレ':'toilet',
                  '給油':'supply'}

        try:
            # 既にあるExcelファイルをロード
            wb = openpyxl.load_workbook('workbook.xlsx')
        except:
            # Excelファイルを新たに生成
            wb = openpyxl.Workbook()
            wb.save('workbook.xlsx')
            ws1 = wb["Sheet"]
            i = 65
            # chr(65)='A'+'1'でA1に列名を挿入　※A->B->C->D...
            for col in column:
                ws1[chr(i) + "1"] = col
                i += 1
            wb.save('workbook.xlsx')

        ws1 = wb["Sheet"]
        j = 65
        # chr(65)='A'+'num'でA[num]に情報を格納  ※A->B->C->D...
        for col in column:
            ws1[chr(j) + str(self.num)] = self.item[column[col]]
            j += 1
        wb.save('workbook.xlsx')

